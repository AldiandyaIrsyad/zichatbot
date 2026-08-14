#!/usr/bin/env python3
"""Replace the obsolete blind-audit claims and Table 4.6 in the current thesis."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(parent: DocumentType):
    for child in parent.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_fragment(paragraph: Paragraph, old: str, new: str) -> bool:
    if old not in paragraph.text:
        return False
    set_paragraph_text(paragraph, paragraph.text.replace(old, new))
    return True


def find_paragraph(doc: DocumentType, starts_with: str) -> Paragraph:
    matches = [p for p in doc.paragraphs if p.text.startswith(starts_with)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting {starts_with!r}, found {len(matches)}")
    return matches[0]


def find_table_after_caption(doc: DocumentType, caption: str) -> Table:
    found = False
    for block in iter_blocks(doc):
        if isinstance(block, Paragraph) and block.text.strip() == caption:
            found = True
            continue
        if found and isinstance(block, Table):
            return block
    raise ValueError(f"No table found after caption {caption!r}")


def percent(value: float) -> str:
    return f"{100 * value:.1f}%".replace(".", ",")


def ci_text(metric: dict) -> str:
    return f"{percent(metric['rate'])} [{percent(metric['ci_low'])}; {percent(metric['ci_high'])}]"


def populate_table(table: Table, summary: dict) -> None:
    while len(table.columns) < 6:
        table.add_column(Inches(0.8))
    while len(table.rows) < 6:
        table.add_row()

    populations = {"subset_a": 150, "subset_b": 160, "subset_c": 200, "subset_d": 210}
    names = {
        "subset_a": "A (validitas QA)",
        "subset_b": "B (keamanan)",
        "subset_c": "C (relevansi)",
        "subset_d": "D (NLI)",
    }
    values = [["Subset", "Pop. (N)", "Audit (n)", "Setuju", "Konkordansi\n(CI Wilson 95%)", "QC benar"]]
    for subset in ("subset_a", "subset_b", "subset_c", "subset_d"):
        authentic = summary["by_subset"][subset]["authentic"]
        qc = summary["by_subset"][subset]["qc"]
        values.append(
            [
                names[subset],
                str(populations[subset]),
                str(authentic["n"]),
                str(authentic["correct"]),
                ci_text(authentic),
                f"{qc['correct']}/{qc['n']}",
            ]
        )
    pooled = summary["pooled_authentic"]
    pooled_qc = summary["pooled_qc"]
    values.append(
        [
            "Total deskriptif",
            "720",
            str(pooled["n"]),
            str(pooled["correct"]),
            ci_text(pooled),
            f"{pooled_qc['correct']}/{pooled_qc['n']}",
        ]
    )

    widths = [1.02, 0.78, 0.82, 0.84, 1.42, 0.74]
    table.autofit = False
    for row_index, row in enumerate(table.rows[:6]):
        for col_index, cell in enumerate(row.cells[:6]):
            cell.text = values[row_index][col_index]
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            for run in paragraph.runs:
                run.font.size = Pt(8)
                run.bold = row_index in (0, 5)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("output_docx", type=Path)
    parser.add_argument(
        "--summary",
        type=Path,
        default=Path("evals/data/blind_check_ad_v2.summary.VERIFIED.json"),
    )
    args = parser.parse_args()

    summary = json.loads(args.summary.read_text(encoding="utf-8"))
    if summary["pooled_authentic"]["n"] != 144 or summary["pooled_authentic"]["correct"] != 141:
        raise ValueError("Unexpected authentic audit totals")
    if summary["pooled_qc"]["n"] != 29 or summary["pooled_qc"]["correct"] != 29:
        raise ValueError("Unexpected QC totals")

    doc = Document(args.input_docx)
    old_id = "pelabelan ulang manusia secara buta terhadap 77 item\u00a0unanimous\u00a0menghasilkan konkordansi 76/77 (98,7%)."
    new_id = "pelabelan ulang manusia secara buta terhadap sampel 20% keempat subset menghasilkan konkordansi 141/144 (97,9%; Wilson 95% CI [94,1%; 99,3%]), sementara 29/29 kontrol QC terdeteksi."
    old_en = "blind human relabelling of 77 unanimous items achieved 76/77 agreement (98.7%)."
    new_en = "blind human relabelling of a 20% sample from all four subsets achieved 141/144 agreement (97.9%; Wilson 95% CI [94.1%, 99.3%]), while all 29 QC controls were detected."
    if sum(replace_fragment(p, old_id, new_id) for p in doc.paragraphs) != 1:
        raise ValueError("Indonesian abstract audit sentence was not replaced exactly once")
    if sum(replace_fragment(p, old_en, new_en) for p in doc.paragraphs) != 1:
        raise ValueError("English abstract audit sentence was not replaced exactly once")

    set_paragraph_text(
        find_paragraph(doc, "Pada dataset v2, pemeriksaan buta diulang"),
        "Pada dataset v2, pemeriksaan buta diulang dengan putusan manusia yang disimpan dan mencakup keempat subset final. Dengan seed 42, sampel terstratifikasi tanpa pengembalian sebesar 20% diambil per subset: A 30/150, B 32/160, C 40/200, dan D 42/210. Sebanyak 29 kontrol QC tambahan disisipkan di luar sampel autentik, sedangkan label referensi, status audit atau QC, dan jawaban tetap tersegel sampai seluruh bagian selesai. Dari 144 baris autentik, 141 konkordan dengan label referensi, yaitu 97,9% (Wilson 95% CI [94,1; 99,3]): A 29/30, B 32/32, C 39/40, dan D 41/42. Seluruh kontrol QC terdeteksi, yaitu 29/29.",
    )
    set_paragraph_text(
        find_paragraph(doc, "Tiga batas mengikat pemeriksaan tersebut."),
        "Empat batas mengikat pemeriksaan baru tersebut. Pertama, pemeriksaan dilakukan oleh satu anotator sehingga tidak menghasilkan inter-rater reliability. Kedua, tugas pelabelan berbeda antar-subset, sehingga total A-D hanya ringkasan deskriptif dan interpretasi utama tetap per subset. Ketiga, kontrol QC adalah konstruksi terkontrol dan bukan baris panel yang benar-benar ditolak. Keempat, meskipun estimasi gabungan 97,9% melampaui 95%, batas bawah Wilson sebesar 94,1% tidak mendukung klaim bahwa konkordansi populasi pasti melebihi 95%. Audit historis 77 item unanimous B/C dipertahankan sebagai hasil terpisah dan tidak digabungkan karena kerangka sampelnya berbeda. Subset D tetap menyimpan lima hasil suara anonim per baris, walaupun tidak menyimpan rasional yang dapat diatribusikan ke model.",
    )
    set_paragraph_text(
        find_paragraph(doc, "Pemeriksaan buta manusia (blind re-labelling)."),
        "Pemeriksaan buta manusia (blind re-labelling). Simpul BLIND pada Gambar 4.2 adalah audit blind injection saat pembuatan dataset. Pemeriksaan yang dilaporkan di sini adalah audit terpisah terhadap dataset v2 final. Dengan seed 42, sebanyak 20% setiap subset diambil secara terstratifikasi tanpa pengembalian, lalu dinilai dalam empat bagian tanpa memperlihatkan label referensi. Kontrol QC tambahan dimasukkan di luar sampel autentik dan dilaporkan terpisah. Prosedur lengkapnya diuraikan pada Bagian 3.1.5.",
    )
    set_paragraph_text(
        find_paragraph(doc, "Konkordansi 98,7%"),
        "Pada 144 baris autentik, 141 label manusia konkordan dengan label referensi, yaitu 97,9% (Wilson 95% CI [94,1; 99,3]). Seluruh 29 kontrol QC terdeteksi dengan benar. Total A-D bersifat deskriptif karena rubrik berbeda antar-subset; hasil utama tetap dibaca per subset. Audit historis B/C sebanyak 77 item unanimous tetap dilaporkan terpisah dan tidak dipool dengan audit baru. Pemeriksaan dilakukan oleh satu anotator sehingga tidak menghasilkan inter-rater reliability.",
    )
    set_paragraph_text(
        find_paragraph(doc, "Dataset bersifat sintetis dengan verifikasi panel model"),
        "Dataset bersifat sintetis dengan verifikasi panel model, bukan gold standard manusia. Blind re-labelling terhadap sampel terstratifikasi 20% dari keempat subset menghasilkan konkordansi 141/144 atau 97,9% (Wilson 95% CI [94,1; 99,3]), sedangkan 29/29 kontrol QC terdeteksi. Pemeriksaan dilakukan oleh satu peneliti sehingga tidak menghasilkan inter-rater reliability dan tetap rentan terhadap kesamaan asumsi antara perancang prosedur dan anotator. Selain itu, total gabungan hanya deskriptif karena rubrik pelabelan berbeda antar-subset, dan batas bawah interval gabungan masih di bawah 95%.",
    )

    table = find_table_after_caption(
        doc, "Tabel 4.6: Hasil Pemeriksaan Buta Manusia (Blind Re-labelling)"
    )
    populate_table(table, summary)

    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output_docx)
    print(f"Wrote {args.output_docx}")


if __name__ == "__main__":
    main()
